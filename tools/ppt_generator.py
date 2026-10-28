from __future__ import annotations

import io
import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt


def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


class PPTGenerator:
    """Builds a .pptx from a SlideOutline and returns its bytes."""

    THEMES = {
        "dark": {"bg": "0C0C0E", "title": "FFFFFF", "body": "C8C8E0", "accent": "7C83FF", "muted": "6A6A82"},
        "light": {"bg": "FFFFFF", "title": "14141F", "body": "33334D", "accent": "4B52D6", "muted": "8A8AA0"},
    }

    BODY_LEFT, BODY_TOP, BODY_W, BODY_H = 0.6, 1.8, 12.1, 5.0
    FONT_SIZES = (20, 18, 16, 14)
    PARA_GAP = 10

    def __init__(self, theme: str = "dark") -> None:
        self.colors = self.THEMES.get(theme, self.THEMES["dark"])

    def generate(self, outline) -> bytes:
        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        slides = [(s.title, list(s.bullets)) for s in outline.slides]
        deck_title = (outline.presentation_title or "").strip() or (slides[0][0] if slides else "")

        # The first slide is a summary under the paper's title. Give the
        # title its own title slide and keep the summary as "Overview".
        if slides and self._similar(slides[0][0], deck_title):
            deck_title = slides[0][0]
            slides[0] = ("Overview", slides[0][1])

        pages: list[tuple[str, list[str]]] = []
        for title, bullets in slides:
            if not bullets:
                continue
            for k, chunk in enumerate(self._paginate(bullets)):
                pages.append((title if k == 0 else f"{title} (cont.)", chunk))

        self._title_slide(prs, deck_title)
        for number, (title, bullets) in enumerate(pages, start=2):
            self._content_slide(prs, title, bullets, number)

        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------
    # Slides
    # ------------------------------------------------------------------

    def _new_slide(self, prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(self.colors["bg"])
        return slide

    def _bar(self, slide, left: float, top: float, width: float) -> None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.07))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(self.colors["accent"])
        bar.line.fill.background()

    def _title_slide(self, prs, title: str) -> None:
        slide = self._new_slide(prs)
        self._bar(slide, 1.0, 2.3, 1.6)
        box = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.3), Inches(2.6))
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.TOP
        p = frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(40 if len(title) <= 80 else 34 if len(title) <= 130 else 28)
        p.font.bold = True
        p.font.color.rgb = _rgb(self.colors["title"])

    def _content_slide(self, prs, title: str, bullets: list[str], number: int) -> None:
        slide = self._new_slide(prs)

        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(1.2))
        box.text_frame.word_wrap = True
        box.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(28 if len(title) <= 70 else 24 if len(title) <= 110 else 20)
        p.font.bold = True
        p.font.color.rgb = _rgb(self.colors["title"])
        self._bar(slide, 0.6, 1.58, 1.0)

        body = slide.shapes.add_textbox(
            Inches(self.BODY_LEFT), Inches(self.BODY_TOP), Inches(self.BODY_W), Inches(self.BODY_H)
        )
        body.text_frame.word_wrap = True
        size = self._pick_size(bullets) or self.FONT_SIZES[-1]
        for j, bullet in enumerate(bullets):
            para = body.text_frame.add_paragraph() if j else body.text_frame.paragraphs[0]
            para.text = f"• {bullet}"
            para.font.size = Pt(size)
            para.font.color.rgb = _rgb(self.colors["body"])
            para.space_before = Pt(self.PARA_GAP)

        num = slide.shapes.add_textbox(Inches(11.9), Inches(6.95), Inches(0.9), Inches(0.35))
        np_ = num.text_frame.paragraphs[0]
        np_.text = str(number)
        np_.font.size = Pt(11)
        np_.font.color.rgb = _rgb(self.colors["muted"])

    # ------------------------------------------------------------------
    # Fit estimation
    # ------------------------------------------------------------------

    def _height_pt(self, bullets: list[str], size: int) -> float:
        chars_per_line = max(20, int((self.BODY_W * 72 - 20) / (size * 0.5)))
        lines = sum(math.ceil((len(b) + 2) / chars_per_line) for b in bullets)
        return lines * size * 1.2 + self.PARA_GAP * len(bullets)

    def _pick_size(self, bullets: list[str]) -> int | None:
        capacity = self.BODY_H * 72 - 10
        return next((s for s in self.FONT_SIZES if self._height_pt(bullets, s) <= capacity), None)

    def _paginate(self, bullets: list[str]) -> list[list[str]]:
        """Split bullets across slides when they would not fit at the smallest size."""
        pages: list[list[str]] = [[]]
        for bullet in bullets:
            if pages[-1] and self._pick_size(pages[-1] + [bullet]) is None:
                pages.append([])
            pages[-1].append(bullet)
        return pages

    @staticmethod
    def _similar(a: str, b: str) -> bool:
        wa, wb = set(a.lower().split()), set(b.lower().split())
        return bool(wa and wb) and len(wa & wb) / min(len(wa), len(wb)) >= 0.8
